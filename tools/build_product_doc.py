from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path


OUT = Path("docs/competition/不平智能求职助手_产品说明文档.docx")
BLUE = "1F4E79"
LIGHT_BLUE = "DCEAF7"
PALE = "F4F7FA"
DARK = "243447"
GRAY = "667085"
GREEN = "31705A"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_header(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    paragraph.add_run(" 页")


def set_run_cn(run, font="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, bold=False, color=None, size=None, align=None, after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_cn(r, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_cn(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_cn(r)
    return p


def add_callout(doc, label, body, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 150, 180, 150, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label)
    set_run_cn(r, bold=True, color=BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(body)
    set_run_cn(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_cn(r, bold=True, color="FFFFFF", size=9.5)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[i], PALE)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(val))
            set_run_cn(r, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def page_break(doc):
    doc.add_page_break()


def title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_run_cn(r, size=22, bold=True, color=BLUE)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(18)
        r2 = p2.add_run(subtitle)
        set_run_cn(r2, size=12, color=GRAY)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after, color in [
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 11.5, 8, 4, DARK),
    ]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(10.5)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_repeat_header(p)
    return p


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    sec.header_distance = Inches(0.4)
    sec.footer_distance = Inches(0.4)
    configure_styles(doc)

    # Cover
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("不平智能求职助手")
    set_run_cn(r, size=30, bold=True, color=BLUE)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(8)
    r2 = p2.add_run("产品说明文档")
    set_run_cn(r2, size=20, bold=True, color=DARK)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(12)
    r3 = p3.add_run("大语言模型驱动的一站式求职辅助平台")
    set_run_cn(r3, size=12, color=GRAY)
    doc.add_paragraph()
    add_callout(doc, "产品定位", "面向高校毕业生、求职者及转岗人群，将简历、岗位、投递与面试信息连接起来，提供从求职准备到进度跟踪的数字化闭环。")
    for _ in range(4):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run("参赛作品说明  |  版本 1.0  |  2026 年 8 月")
    set_run_cn(rm, size=10, color=GRAY)
    link = doc.add_paragraph()
    link.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rl = link.add_run("项目地址：https://github.com/yi-wang-2/Buping_Job_Seeker_Assistant/tree/dev")
    set_run_cn(rl, size=9, color=BLUE)

    # headers footers
    header = sec.header.paragraphs[0]
    header.text = "不平智能求职助手｜产品说明文档"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_cn(run, size=8.5, color=GRAY)
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        set_run_cn(run, size=8.5, color=GRAY)

    page_break(doc)
    title(doc, "文档概览", "本说明书用于介绍参赛产品的定位、功能、技术实现、创新价值与应用前景。")
    add_table(doc, ["项目", "说明"], [
        ["产品名称", "不平智能求职助手（Buping Job Seeker Assistant）"],
        ["产品形态", "Web 全栈应用，可在本地或服务器环境部署"],
        ["目标用户", "高校毕业生、社会求职者、转岗人群及职业辅导人员"],
        ["核心场景", "简历生成与优化、岗位分析与推荐、求职进度管理、模拟面试"],
        ["核心技术", "Python、FastAPI、Pydantic v2、LangChain、React 19、TypeScript、Vite 6"],
        ["文档依据", "基于项目 dev 分支现有源码、依赖清单与 README 整理"],
    ], [1.45, 5.05])
    h(doc, "目录", 1)
    for item in [
        "1. 产品背景与问题定义", "2. 产品定位与目标用户", "3. 产品总体方案",
        "4. 核心功能说明", "5. 典型使用流程", "6. 技术架构与实现",
        "7. 产品创新点", "8. 安全、隐私与可靠性", "9. 应用价值与竞争优势",
        "10. 部署运行与未来规划", "附录：技术栈与产品边界",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "说明", "文档中的“约 30 秒”等耗时为典型环境下的产品设计参考，实际结果会随模型服务、网络条件、输入长度和硬件环境变化。")

    page_break(doc)
    title(doc, "1. 产品背景与问题定义")
    h(doc, "1.1 行业背景", 2)
    add_text(doc, "求职不是单点任务，而是一条由信息收集、岗位筛选、简历定制、投递记录、面试准备和结果复盘构成的连续流程。现有工具往往只解决其中一个环节，用户需要在招聘平台、文档工具、表格和聊天机器人之间反复切换，信息分散且难以沉淀。")
    h(doc, "1.2 用户痛点", 2)
    add_table(doc, ["痛点", "具体表现", "产品响应"], [
        ["简历制作门槛高", "内容组织、表达优化和版式设计耗时", "AI 生成、岗位定制、可视化编辑与模板导出"],
        ["岗位信息分散", "岗位表、网页和个人记录难以统一管理", "岗位雷达导入、评分推荐与收藏/忽略操作"],
        ["投递过程不可见", "投递后容易遗忘时间、状态和后续动作", "求职看板、状态跟进、定时检查与通知"],
        ["面试准备碎片化", "题库缺少针对性，难以形成真实对话压力", "基于简历与 JD 的准备报告和多轮模拟面试"],
        ["AI 使用缺乏连续性", "不同任务重复输入背景，上下文成本高", "结构化配置、记忆仓库、缓存与上下文预算"],
    ], [1.2, 2.4, 2.9])
    h(doc, "1.3 产品机会", 2)
    add_text(doc, "不平智能求职助手的核心机会，是把大语言模型从“单次问答工具”转化为可持续使用的求职工作台：以个人简历和求职偏好为基础，以岗位描述为任务上下文，以投递记录和面试反馈为过程数据，形成可复用、可追踪、可迭代的求职资产。")

    page_break(doc)
    title(doc, "2. 产品定位与目标用户")
    h(doc, "2.1 产品定位", 2)
    add_callout(doc, "一句话定义", "不平智能求职助手是一款由大语言模型驱动的全栈式求职辅助平台，为用户提供“简历—岗位—投递—面试—复盘”的一站式工作流。")
    h(doc, "2.2 目标用户", 2)
    add_table(doc, ["用户类型", "主要需求", "高频功能"], [
        ["高校毕业生", "快速建立专业简历，理解岗位要求，准备首次面试", "简历生成、JD 分析、面试准备"],
        ["社会求职者", "针对不同岗位高效改写简历，管理多条投递", "岗位定制、求职看板、状态跟进"],
        ["跨行业转岗者", "识别能力差距，重构经历表达，补充面试论证", "技能匹配、职业建议、模拟面试"],
        ["职业辅导人员", "辅助学员形成标准化材料并开展训练", "简历校验、报告导出、面试评估"],
    ], [1.35, 2.9, 2.25])
    h(doc, "2.3 产品目标", 2)
    for text in [
        "降低求职材料制作门槛，让用户从零散信息快速形成结构化、ATS 友好的专业简历。",
        "提高岗位理解和简历定制效率，使简历内容与岗位要求建立清晰对应关系。",
        "减少投递信息遗漏，通过统一看板、状态跟进和通知提升求职过程可控性。",
        "建立可反复训练的面试环境，为用户提供个性化提问、追问与多维评估。",
        "在可配置模型和本地数据管理基础上，为个人求职数据提供更灵活的控制方式。",
    ]:
        add_bullet(doc, text)

    page_break(doc)
    title(doc, "3. 产品总体方案")
    h(doc, "3.1 产品闭环", 2)
    add_number(doc, "建立个人基础：配置模型服务、导入历史简历、维护个人信息与求职偏好。")
    add_number(doc, "获取岗位机会：从文件、复制表格或在线岗位表导入岗位，按个人画像计算匹配度。")
    add_number(doc, "生成求职材料：解析简历，结合岗位描述生成或优化 ATS 友好简历，并完成在线编辑和导出。")
    add_number(doc, "管理投递过程：记录公司、岗位、状态、标签、时间和备注，按需执行状态跟进。")
    add_number(doc, "开展面试训练：生成准备报告，选择面试官风格进行多轮对话，结束后形成评估报告。")
    add_number(doc, "沉淀过程资产：保存岗位分析、AI 使用指标、记忆与历史结果，支持后续复用和复盘。")
    h(doc, "3.2 功能地图", 2)
    add_table(doc, ["产品域", "主要能力", "输出结果"], [
        ["简历中心", "多格式解析、生成、定制、校验、编辑、预览、导出", "结构化简历、HTML 预览、PDF 文件"],
        ["岗位雷达", "岗位导入、偏好配置、匹配评分、每日推荐", "候选岗位列表与匹配说明"],
        ["AI 技能", "JD 分析、技能匹配、职业建议、文本改写", "结构化分析与行动建议"],
        ["求职看板", "投递记录、标签、状态、统计、自动跟进", "可追踪的求职进度数据"],
        ["面试中心", "面试准备、模拟对话、语音播报、综合评估", "准备报告、对话记录、评估报告"],
        ["系统能力", "模型配置、记忆、缓存、链路追踪、通知", "稳定且可配置的 AI 运行环境"],
    ], [1.2, 3.5, 1.8])

    page_break(doc)
    title(doc, "4. 核心功能说明（上）")
    h(doc, "4.1 智能简历解析与结构化", 2)
    add_text(doc, "系统采用分层解析策略：YAML/JSON 等结构化文件直接读取；PDF、DOCX、HTML、Markdown、TXT、LaTeX 等文档先提取文本，再由大语言模型映射为统一简历结构。PDF 提取优先使用 PyMuPDF，并以 pdfminer 作为回退方案。")
    add_bullet(doc, "覆盖 8 类常见文档类型：YAML、JSON、TXT、Markdown、PDF、DOCX、HTML、LaTeX。")
    add_bullet(doc, "统一抽取个人信息、教育经历、工作经历、项目、成果、证书、语言和兴趣等字段。")
    add_bullet(doc, "提供字段完整性校验和错误/警告提示，减少空字段或结构异常进入生成流程。")
    h(doc, "4.2 AI 简历生成与岗位定制", 2)
    add_text(doc, "系统将用户基础资料、目标语言和岗位描述组织为模型上下文，生成结构完整、表达专业的 ATS 友好简历。普通生成采用一次核心 LLM 调用；岗位定制根据 JD 强调匹配经历、技能与量化成果。")
    add_callout(doc, "效率参考", "典型网络与模型服务条件下，普通简历生成约 30 秒。该时间为经验值，不构成固定性能承诺。", fill="EAF3EE")
    h(doc, "4.3 可视化编辑与多样式输出", 2)
    add_bullet(doc, "WYSIWYG 编辑：支持标题、粗体、斜体、下划线、列表、引用、链接、撤销/重做等操作。")
    add_bullet(doc, "布局微调：支持正文行距、模块间距和内容块编辑，并提供自动保存与重置能力。")
    add_bullet(doc, "实时预览与导出：将结构化简历渲染为 HTML，并通过浏览器渲染链路生成 PDF。")
    add_bullet(doc, "样式系统：内置 5 套专业 CSS 简历模板，可在内容不变的情况下切换视觉风格。")
    h(doc, "4.4 岗位分析与能力匹配", 2)
    add_text(doc, "AI 技能模块提供岗位描述分析、技能匹配、职业建议与文本改写能力。岗位分析结果可归档，便于用户比较不同职位的职责、硬性要求、关键词和潜在能力缺口。")

    page_break(doc)
    title(doc, "4. 核心功能说明（下）")
    h(doc, "4.5 岗位雷达与推荐", 2)
    add_text(doc, "岗位雷达可导入 CSV/表格类岗位数据、粘贴的表格文本，或同步受支持的在线岗位表。系统结合用户技能、期望地点、岗位方向等偏好计算匹配分数，并提供推荐、收藏和忽略等操作。")
    h(doc, "4.6 求职看板与状态跟进", 2)
    add_bullet(doc, "集中记录公司、岗位、链接、当前状态、标签、备注和关键时间。")
    add_bullet(doc, "提供统计概览，帮助用户识别投递、面试和结果阶段的数量分布。")
    add_bullet(doc, "在用户完成目标平台登录连接后，可按需或按计划检查投递状态，并保留识别证据和置信信息。")
    add_bullet(doc, "支持邮件和微信类通知渠道配置，用于发送状态变化或测试消息。")
    h(doc, "4.7 面试准备与多轮模拟", 2)
    add_text(doc, "面试模块依据个人简历、目标岗位和公司信息生成针对性准备内容。模拟面试支持友善型、专业型、压力型、学术型和闲聊型 5 种面试官风格，根据候选人最新回答继续追问或切换话题。")
    add_table(doc, ["能力", "实现说明"], [
        ["多轮上下文", "保留近期对话并进行上下文预算控制，避免问题重复和无关扩展"],
        ["轮次管理", "覆盖开场、项目深挖、技术、行为、反问和结束等阶段"],
        ["结果评估", "从技术能力、沟通表达、综合素质、优势与改进方向等维度输出报告"],
        ["语音能力", "支持多种 TTS 实现，为面试官文本提供语音播报或流式输出"],
    ], [1.35, 5.15])
    h(doc, "4.8 AI 工程化支撑", 2)
    add_text(doc, "项目实现统一模型网关、上下文预算、增量上下文、提示词缓存、记忆仓库、调用追踪与指标统计等能力，使不同业务模块能够复用一致的 AI 调用机制，并提升可观察性和可维护性。")

    page_break(doc)
    title(doc, "5. 典型使用流程")
    h(doc, "5.1 场景一：针对目标岗位生成定制简历", 2)
    add_number(doc, "用户在设置页配置可用的大模型服务，并导入已有简历或录入个人信息。")
    add_number(doc, "系统解析文档并转换为统一结构，执行必填字段和内容完整性校验。")
    add_number(doc, "用户粘贴目标岗位 JD，选择简历语言和视觉模板。")
    add_number(doc, "AI 识别岗位关键词与能力要求，生成或重写与岗位更匹配的简历内容。")
    add_number(doc, "用户在可视化编辑器中调整文本和版式，实时预览并导出 PDF。")
    h(doc, "5.2 场景二：从岗位发现到投递跟进", 2)
    add_number(doc, "用户导入岗位表并配置目标职位、地点、技能等偏好。")
    add_number(doc, "岗位雷达计算匹配度，用户查看每日推荐并筛选目标岗位。")
    add_number(doc, "选定岗位后进入简历定制与投递准备，将结果加入求职看板。")
    add_number(doc, "用户持续更新投递状态，或连接目标平台后执行状态检查。")
    add_number(doc, "系统在状态发生变化时更新记录，并根据配置发送提醒。")
    h(doc, "5.3 场景三：模拟面试与复盘", 2)
    add_number(doc, "选择目标岗位、公司、面试类型和面试官风格。")
    add_number(doc, "系统读取简历与岗位上下文，由 AI 面试官生成开场问题。")
    add_number(doc, "用户逐轮回答，系统基于回答内容继续追问并推进面试阶段。")
    add_number(doc, "结束后生成综合评估报告，用户据此改进答案并再次训练。")

    page_break(doc)
    title(doc, "6. 技术架构与实现")
    h(doc, "6.1 总体架构", 2)
    add_table(doc, ["层级", "主要组件", "职责"], [
        ["表现层", "React 19、TypeScript、Tailwind CSS、TipTap/iframe 编辑器", "页面交互、富文本编辑、数据展示与主题适配"],
        ["接口层", "FastAPI、Pydantic v2", "REST API、参数校验、文件上传与静态资源托管"],
        ["业务层", "简历、岗位雷达、求职跟进、面试、通知等服务", "组织业务规则和跨模块工作流"],
        ["AI 能力层", "LangChain、统一 LLM Gateway、技能注册与运行时", "模型调用、技能编排、错误处理与调用追踪"],
        ["上下文层", "Token Budget、Context Manager、增量上下文、缓存", "控制输入规模，复用结果并优化调用成本"],
        ["数据层", "SQLite、JSON/YAML 文件、模型与输出目录", "保存配置、记忆、指标、岗位与求职记录"],
        ["文档/媒体层", "PyMuPDF、pdfminer、python-docx、BeautifulSoup、ReportLab、TTS", "文档抽取、PDF 生成和语音合成"],
    ], [1.1, 2.65, 2.75])
    h(doc, "6.2 前后端协作", 2)
    add_text(doc, "前端通过统一 API 客户端调用 FastAPI 服务。开发环境下前后端独立运行，生产构建后可由后端托管前端静态资源。接口使用 Pydantic 模型进行请求和响应约束，业务服务负责文件、模型、数据库及外部页面的具体操作。")
    h(doc, "6.3 可扩展模型接入", 2)
    add_text(doc, "模型层将供应商、模型名称、基础地址和鉴权信息抽象为可配置项，兼容 Anthropic 风格和 OpenAI 风格的模型接口。业务模块通过统一网关调用模型，减少直接依赖单一厂商的耦合。")

    page_break(doc)
    title(doc, "7. 产品创新点")
    h(doc, "7.1 从单点 AI 功能升级为求职闭环", 2)
    add_text(doc, "产品并非只提供“生成一份简历”或“回答一道面试题”，而是围绕个人画像连接岗位发现、材料定制、投递跟踪和面试复盘，使 AI 输出能够进入后续业务流程并持续产生价值。")
    h(doc, "7.2 结构化数据与生成式 AI 协同", 2)
    add_text(doc, "简历、岗位、投递和面试信息均以结构化对象进入系统。大语言模型负责理解、改写与推理，规则校验、状态机和数据库负责确定性约束，从而降低完全依赖自由文本生成带来的不稳定性。")
    h(doc, "7.3 面向个人场景的 AI 工程化", 2)
    add_text(doc, "项目引入统一网关、上下文预算、缓存、记忆和观测指标，将大模型调用从页面级功能提升为平台能力。其价值不仅是节省 Token，也包括统一故障处理、追踪调用路径和持续评估生成效果。")
    h(doc, "7.4 多风格、可追问的模拟面试", 2)
    add_text(doc, "模拟面试不是固定题库顺序播放，而是结合简历、JD、面试阶段和最近回答动态生成下一问。5 种面试官风格让用户能够在不同沟通压力和考察角度下反复训练。")
    h(doc, "7.5 可控的数据与部署方式", 2)
    add_text(doc, "平台支持本地运行和自定义模型服务地址，用户可以依据场景选择云端或兼容接口的模型服务。求职记录、记忆和生成文件可保存在本地项目数据目录，便于个人控制与迁移。")
    add_callout(doc, "创新总结", "以“业务闭环 + 结构化数据 + 生成式 AI + 工程化治理”为核心，将多个高频求职任务组合为统一的个人求职操作系统。", fill="EAF3EE")

    page_break(doc)
    title(doc, "8. 安全、隐私与可靠性")
    h(doc, "8.1 数据与密钥管理", 2)
    add_bullet(doc, "模型密钥和服务地址通过配置管理，前端设置接口对敏感信息进行受控读写。")
    add_bullet(doc, "简历、求职记录、记忆、指标和输出文件主要保存在本地数据目录或本地 SQLite 数据库。")
    add_bullet(doc, "产品允许用户自行选择模型服务；是否向第三方模型发送内容取决于用户配置和所使用的服务。")
    h(doc, "8.2 输入与输出可靠性", 2)
    add_bullet(doc, "使用 Pydantic 和业务校验器约束接口参数与简历数据结构。")
    add_bullet(doc, "PDF 文本抽取采用主方案与回退方案，异常情况下记录诊断信息。")
    add_bullet(doc, "LLM 输出经过候选内容提取、结构解析和默认骨架补全，减少格式漂移导致的流程中断。")
    add_bullet(doc, "模型网关提供超时、重试、错误映射和追踪能力，关键调用能够留存指标。")
    h(doc, "8.3 自动跟进边界", 2)
    add_text(doc, "投递状态跟进需要用户主动连接并完成目标平台登录。状态识别结合页面证据、规则和可选的模型判断，并记录置信信息。由于招聘平台页面可能变化，产品将该能力定位为辅助检查，最终状态应由用户核验。")
    h(doc, "8.4 使用提示", 2)
    add_callout(doc, "隐私提示", "简历通常包含姓名、电话、邮箱和工作经历等敏感信息。比赛演示建议使用脱敏数据；实际使用时，应选择可信模型服务并妥善保管 API Key。", fill="FFF4E5")

    page_break(doc)
    title(doc, "9. 应用价值与竞争优势")
    h(doc, "9.1 用户价值", 2)
    add_table(doc, ["价值维度", "价值体现"], [
        ["效率", "减少简历重写、岗位整理、投递记录和面试材料准备中的重复劳动"],
        ["质量", "通过岗位上下文、结构校验和专业模板提升求职材料的一致性与可读性"],
        ["连续性", "将分散任务纳入同一工作流，保留可复用的个人资料和过程数据"],
        ["可训练性", "通过多轮模拟与评估支持反复练习，而非一次性获取答案"],
        ["可控性", "支持自定义模型接口和本地数据存储，便于个人化配置与部署"],
    ], [1.3, 5.2])
    h(doc, "9.2 差异化优势", 2)
    add_table(doc, ["对比维度", "通用聊天机器人", "单点简历工具", "不平智能求职助手"], [
        ["业务范围", "依赖用户逐次提问", "聚焦简历生成", "覆盖简历、岗位、投递、面试与复盘"],
        ["数据组织", "以对话文本为主", "通常为表单或文档", "统一结构化个人、岗位与过程数据"],
        ["流程衔接", "输出需手动搬运", "导出后流程结束", "AI 输出可进入看板、定制和训练流程"],
        ["面试能力", "可问答但流程松散", "通常不包含", "多风格、多轮追问、轮次管理与评估"],
        ["部署控制", "由平台统一托管", "取决于供应商", "支持本地运行及兼容模型接口配置"],
    ], [1.1, 1.65, 1.65, 2.1])
    h(doc, "9.3 适用场景", 2)
    add_text(doc, "产品可用于个人求职准备、高校就业辅导、职业转型训练、求职训练营演示及 AI 应用教学。对于机构化场景，后续可在权限、团队协作和数据隔离等方面继续扩展。")

    page_break(doc)
    title(doc, "10. 部署运行与未来规划")
    h(doc, "10.1 运行环境", 2)
    add_table(doc, ["项目", "要求/说明"], [
        ["后端", "Python 3.10—3.12；FastAPI 应用与相关依赖"],
        ["前端", "Node.js 构建环境；React 19、TypeScript、Vite 6"],
        ["模型服务", "可用的 Anthropic 风格或 OpenAI 风格兼容接口及 API Key"],
        ["浏览器能力", "简历 PDF 渲染及部分状态跟进功能需要可用浏览器环境"],
        ["语音能力", "按所选 TTS 方案安装模型或配置对应服务"],
        ["数据存储", "本地文件目录与 SQLite；部署时应设置持久化存储和备份策略"],
    ], [1.4, 5.1])
    h(doc, "10.2 基本部署流程", 2)
    add_number(doc, "准备 Python 环境并安装项目依赖。")
    add_number(doc, "进入 frontend 目录安装前端依赖并执行构建。")
    add_number(doc, "配置模型服务、API Key、数据目录及可选通知/TTS 参数。")
    add_number(doc, "启动 FastAPI 服务，由后端提供 API 并托管已构建的前端资源。")
    add_number(doc, "使用演示或脱敏数据验证简历生成、岗位导入和模拟面试主流程。")
    h(doc, "10.3 后续规划", 2)
    for text in [
        "增强岗位来源接入能力，在合法合规前提下支持更多标准化数据源。",
        "建立更完整的简历质量、岗位匹配和面试表现评测集，持续验证 AI 输出质量。",
        "完善账号、权限、数据加密与多租户隔离，为机构化部署提供基础。",
        "增加任务队列和异步处理能力，改善长耗时模型调用与文档生成体验。",
        "扩展可解释推荐与求职分析看板，帮助用户发现投递策略和能力短板。",
        "完善移动端适配、无障碍体验和更多语言支持。",
    ]:
        add_bullet(doc, text)

    page_break(doc)
    title(doc, "附录 A：主要技术栈")
    add_table(doc, ["类别", "技术/组件", "用途"], [
        ["后端框架", "Python、FastAPI、Pydantic v2、Uvicorn", "API、数据模型与服务运行"],
        ["前端框架", "React 19、TypeScript、Vite 6、Tailwind CSS", "Web 界面与工程构建"],
        ["AI 编排", "LangChain、LangSmith、自研 LLM Gateway", "模型调用、技能运行与追踪"],
        ["模型适配", "OpenAI/Anthropic 风格兼容接口", "支持可配置模型服务"],
        ["文档处理", "PyMuPDF、pdfminer.six、python-docx、BeautifulSoup、lxml", "文档抽取与解析"],
        ["文档输出", "ReportLab、Selenium/Chrome 渲染", "报告及简历 PDF 输出"],
        ["编辑器", "TipTap、iframe designMode", "富文本和整页简历编辑"],
        ["语音", "edge-tts、Kokoro、ChatTTS、MiniMax TTS 接口", "模拟面试语音播报"],
        ["存储", "SQLite、JSON、YAML、本地文件", "业务记录、配置与生成物"],
        ["测试", "pytest、pytest-mock、pytest-cov", "单元、集成与质量验证"],
    ], [1.2, 2.7, 2.6])
    h(doc, "附录 B：产品边界与声明", 1)
    add_bullet(doc, "本产品是求职辅助工具，不保证用户获得面试或录用结果。")
    add_bullet(doc, "AI 生成内容可能存在遗漏或不准确，用户在投递前应核实个人经历、数据和事实。")
    add_bullet(doc, "岗位匹配分数用于排序和辅助判断，不替代用户对岗位、公司与职业方向的自主决策。")
    add_bullet(doc, "平台页面结构和访问策略可能变化，自动状态跟进功能应以实际运行结果和人工核验为准。")
    add_bullet(doc, "项目未使用 OpenCV；图像/文档能力主要由浏览器渲染、PDF 与文档处理库实现。")
    add_callout(doc, "结语", "不平智能求职助手以真实求职流程为牵引，将生成式 AI 的理解与表达能力嵌入可操作、可追踪的产品闭环，为个人提供更高效、更系统、更可控的求职体验。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "不平智能求职助手 产品说明文档"
    doc.core_properties.subject = "参赛作品产品说明"
    doc.core_properties.author = "不平智能求职助手项目组"
    doc.core_properties.keywords = "AI求职,简历,岗位匹配,模拟面试,FastAPI,React,LangChain"
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
