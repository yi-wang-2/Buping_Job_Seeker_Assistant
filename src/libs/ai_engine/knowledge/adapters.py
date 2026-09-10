from __future__ import annotations

import hashlib
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

from .models import KnowledgeSource, RawKnowledgeUnit, SourceDocument


SUPPORTED_SUFFIXES = {".md", ".markdown", ".json", ".pdf", ".docx", ".zip"}
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024


def _hash(content: str) -> str:
    return hashlib.sha256(content.encode("utf-8")).hexdigest()


class LocalFileAdapter:
    def discover(self, source: KnowledgeSource) -> list[SourceDocument]:
        location = Path(source.location).resolve()
        paths = [location] if location.is_file() else sorted(
            path for path in location.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
        )
        documents = []
        for path in paths:
            if path.suffix.lower() not in SUPPORTED_SUFFIXES or path.stat().st_size > MAX_DOCUMENT_BYTES:
                continue
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                content = self._pdf_text(path)
                media_type = "application/pdf"
            elif suffix == ".docx":
                content = self._docx_text(path)
                media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif suffix == ".zip":
                content = ""
                media_type = "application/zip"
            else:
                content = path.read_text(encoding="utf-8-sig")
                media_type = "application/json" if suffix == ".json" else "text/markdown"
            documents.append(SourceDocument(
                id=hashlib.sha256(f"{source.id}:{path}".encode()).hexdigest()[:24],
                source_id=source.id,
                path=str(path),
                media_type=media_type,
                content=content,
                content_hash=_hash(content),
                metadata={"relative_path": path.name if location.is_file() else path.relative_to(location).as_posix()},
            ))
        return documents

    def parse(self, document: SourceDocument) -> list[RawKnowledgeUnit]:
        if document.media_type == "application/json":
            return self._parse_json(document)
        if document.media_type == "application/zip":
            return self._parse_archive(document)
        return self._parse_markdown(document)

    def fingerprint(self, document: SourceDocument) -> str:
        return document.content_hash

    @staticmethod
    def _pdf_text(path: Path) -> str:
        from pdfminer.high_level import extract_text
        return extract_text(str(path)).strip()

    @staticmethod
    def _docx_text(path: Path) -> str:
        from docx import Document
        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())

    @classmethod
    def _parse_archive(cls, document: SourceDocument) -> list[RawKnowledgeUnit]:
        units: list[RawKnowledgeUnit] = []
        archive_path = Path(document.path)
        with zipfile.ZipFile(archive_path) as archive:
            for info in archive.infolist():
                if info.is_dir() or info.file_size > MAX_DOCUMENT_BYTES:
                    continue
                member = Path(info.filename)
                if member.is_absolute() or ".." in member.parts:
                    continue
                suffix = member.suffix.lower()
                if suffix not in {".md", ".markdown", ".json"}:
                    continue
                content = archive.read(info).decode("utf-8-sig")
                nested = SourceDocument(
                    id=hashlib.sha256(f"{document.id}:{info.filename}".encode()).hexdigest()[:24],
                    source_id=document.source_id, path=f"{document.path}!/{info.filename}",
                    media_type="application/json" if suffix == ".json" else "text/markdown",
                    content=content, content_hash=_hash(content),
                    metadata={"archive": document.path, "relative_path": info.filename},
                )
                units.extend(cls._parse_json(nested) if suffix == ".json" else cls._parse_markdown(nested))
        return units

    @staticmethod
    def _parse_markdown(document: SourceDocument) -> list[RawKnowledgeUnit]:
        sections: list[tuple[str, list[str]]] = []
        title = Path(document.path).stem
        body: list[str] = []
        for line in document.content.splitlines():
            match = re.match(r"^#{1,4}\s+(.+?)\s*$", line)
            if match:
                if any(item.strip() for item in body):
                    sections.append((title, body))
                title, body = match.group(1).strip(), []
            else:
                body.append(line)
        if any(item.strip() for item in body):
            sections.append((title, body))
        return [RawKnowledgeUnit(
            unit_type="concept", title=heading, content="\n".join(lines).strip(),
            source_path=f"{document.path}#{index + 1}",
            metadata={"document_id": document.id, "heading": heading},
        ) for index, (heading, lines) in enumerate(sections) if "\n".join(lines).strip()]

    @classmethod
    def _parse_json(cls, document: SourceDocument) -> list[RawKnowledgeUnit]:
        payload = json.loads(document.content)
        units: list[RawKnowledgeUnit] = []

        def walk(value: Any, path: str, inherited: dict[str, str]) -> None:
            if isinstance(value, list):
                for index, item in enumerate(value):
                    walk(item, f"{path}/{index}", inherited)
                return
            if not isinstance(value, dict):
                return
            context = dict(inherited)
            for key in ("company", "role", "topic", "category", "difficulty"):
                if value.get(key):
                    context[key] = str(value[key])
            question = value.get("question") or value.get("题目") or value.get("问题")
            answer = value.get("answer") or value.get("reference_answer") or value.get("答案")
            if question and answer:
                intent = value.get("interviewer_intent") or value.get("thinking") or value.get("考察意图") or "未标注"
                competencies = value.get("competencies") or value.get("能力") or value.get("考察点") or []
                if isinstance(competencies, str):
                    competencies = [competencies]
                units.append(RawKnowledgeUnit(
                    unit_type="question_card", title=str(question)[:500],
                    content=f"问题：{question}\n参考答案：{answer}", source_path=f"{document.path}#{path}",
                    metadata={
                        **context, "question": str(question), "reference_answer": str(answer),
                        "competencies": list(competencies), "interviewer_intent": str(intent),
                        "follow_ups": value.get("follow_ups") or value.get("追问") or [],
                        "scoring_points": value.get("scoring_points") or value.get("评分要点") or [],
                        "document_id": document.id,
                    },
                ))
                return
            for key, item in value.items():
                walk(item, f"{path}/{key}", context)

        walk(payload, "$", {})
        return units


class MarkdownAdapter(LocalFileAdapter):
    pass


class JSONAdapter(LocalFileAdapter):
    pass


class PDFAdapter(LocalFileAdapter):
    pass


class DocxAdapter(LocalFileAdapter):
    pass


class ArchiveAdapter(LocalFileAdapter):
    pass


class UploadAdapter(LocalFileAdapter):
    pass


class GitSourceAdapter(LocalFileAdapter):
    """Read an already checked-out Git working tree without invoking Git or network access."""

    EXCLUDED_PARTS = {".git", ".github", "agents", "scripts", "node_modules", "dist", "build"}

    def discover(self, source: KnowledgeSource) -> list[SourceDocument]:
        documents = super().discover(source)
        return [
            document for document in documents
            if not self.EXCLUDED_PARTS.intersection(
                Path(document.metadata.get("relative_path") or document.path).parts
            )
        ]
